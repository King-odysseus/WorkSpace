import json
import csv
import io
import mimetypes
import os
import base64
import hashlib
import logging
from urllib import request as urlrequest
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse

from cryptography.fernet import Fernet, InvalidToken
from django.conf import settings as django_settings
from django.db import transaction
from django.http import JsonResponse
from django.http import HttpResponse, HttpResponseRedirect
from openpyxl import Workbook, load_workbook
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from .models import AiAction, Membership, WorkspaceDocument, WorkspaceDocumentComment, WorkspaceDocumentRevision, WorkspaceDocumentShare, WorkspaceFile, WorkspaceSetting
from .file_responses import stored_file_response
from .sanitize import sanitize_document_content
from .views import parse_int, require_workspace_member
from .ai_actions import (
    ActionExecutionError,
    ActionValidationError,
    PrivacyBoundaryError,
    PrivacyRegistry,
    action_instructions,
    build_workspace_snapshot,
    create_action_proposal,
    execute_action,
    parse_provider_response,
)

logger = logging.getLogger(__name__)


def _setting(workspace_id):
    return WorkspaceSetting.objects.get_or_create(workspace_id=workspace_id)[0]


def _document_permission(document, membership, user):
    if membership.role in {'owner', 'manager'} or document.created_by_id == user.id:
        return 'edit'
    share = document.shares.filter(user=user).only('permission').first()
    return share.permission if share else 'view'


# Autosave fires every time typing pauses, so an unthrottled snapshot per change
# would store dozens of full document copies per editing session.
REVISION_THROTTLE_SECONDS = 300
REVISION_KEEP = 40


def _record_revision(document, user):
    """Snapshot the pre-edit content, collapsing rapid autosaves by one author."""
    latest = document.revisions.first()
    if latest and latest.created_by_id == user.id and (timezone.now() - latest.created_at).total_seconds() < REVISION_THROTTLE_SECONDS:
        return
    WorkspaceDocumentRevision.objects.create(document=document, created_by=user, title=document.title, content=document.content or {})
    stale = list(document.revisions.values_list('id', flat=True)[REVISION_KEEP:])
    if stale:
        WorkspaceDocumentRevision.objects.filter(id__in=stale).delete()


AI_PROVIDER_DEFAULTS = {
    'openai': {'base_url': 'https://api.openai.com/v1', 'model': 'gpt-4o-mini', 'key_env': ('OPENAI_API_KEY', 'AI_API_KEY'), 'url_env': 'AI_API_URL', 'model_env': 'AI_MODEL'},
    'claude': {'base_url': 'https://api.anthropic.com/v1', 'model': 'claude-3-5-haiku-latest', 'key_env': ('ANTHROPIC_API_KEY',), 'url_env': 'CLAUDE_API_URL', 'model_env': 'CLAUDE_MODEL'},
    'kimi': {'base_url': 'https://api.moonshot.cn/v1', 'model': 'moonshot-v1-8k', 'key_env': ('KIMI_API_KEY',), 'url_env': 'KIMI_API_URL', 'model_env': 'KIMI_MODEL'},
    'deepseek': {'base_url': 'https://api.deepseek.com', 'model': 'deepseek-v4-flash', 'key_env': ('DEEPSEEK_API_KEY',), 'url_env': 'DEEPSEEK_API_URL', 'model_env': 'DEEPSEEK_MODEL'},
}


def _secret_cipher():
    key = base64.urlsafe_b64encode(hashlib.sha256(django_settings.SECRET_KEY.encode('utf-8')).digest())
    return Fernet(key)


def _encrypt_secret(value):
    return _secret_cipher().encrypt(value.encode('utf-8')).decode('ascii')


def _decrypt_secret(value):
    if not value:
        return ''
    try:
        return _secret_cipher().decrypt(value.encode('ascii')).decode('utf-8')
    except (InvalidToken, ValueError, TypeError):
        return ''


def _environment_value(names):
    return next((os.environ.get(name, '').strip() for name in names if os.environ.get(name, '').strip()), '')


def _provider_values(setting, provider):
    defaults = AI_PROVIDER_DEFAULTS[provider]
    stored = (setting.ai_provider_config or {}).get(provider, {})
    api_key = _decrypt_secret(stored.get('api_key_encrypted')) or _environment_value(defaults['key_env'])
    base_url = stored.get('base_url') or os.environ.get(defaults['url_env'], '').strip() or defaults['base_url']
    model = stored.get('model') or os.environ.get(defaults['model_env'], '').strip() or (setting.ai_model if provider == setting.ai_default_provider else '') or defaults['model']
    return {'api_key': api_key, 'base_url': base_url.rstrip('/'), 'model': model}


def _safe_provider_config(setting):
    result = {}
    for provider in AI_PROVIDER_DEFAULTS:
        values = _provider_values(setting, provider)
        result[provider] = {
            'base_url': values['base_url'],
            'model': values['model'],
            'has_api_key': bool(values['api_key']),
            'key_hint': f"••••{values['api_key'][-4:]}" if values['api_key'] else '',
        }
    return result


def _provider_endpoint(provider, base_url):
    base_url = base_url.rstrip('/')
    expected_path = '/messages' if provider == 'claude' else '/chat/completions'
    return base_url if base_url.endswith(expected_path) else f'{base_url}{expected_path}'


@require_http_methods(['GET', 'PATCH'])
def workspace_check_in_settings(request, workspace_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    setting = _setting(workspace_id)
    if request.method == 'GET':
        return JsonResponse({'settings': setting.as_dict(), 'can_manage': membership.role in {'owner', 'manager'}})
    if membership.role not in {'owner', 'manager'}:
        return JsonResponse({'error': 'Owner or manager access is required.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    if 'check_in_reminder_hour' not in payload:
        return JsonResponse({'error': 'Reminder hour is required.'}, status=400)
    try:
        reminder_hour = int(payload['check_in_reminder_hour'])
    except (TypeError, ValueError):
        return JsonResponse({'error': 'Reminder hour must be an integer from 0 to 23.'}, status=400)
    if not 0 <= reminder_hour <= 23:
        return JsonResponse({'error': 'Reminder hour must be an integer from 0 to 23.'}, status=400)
    setting.check_in_reminder_hour = reminder_hour
    setting.save(update_fields=['check_in_reminder_hour', 'updated_at'])
    return JsonResponse({'settings': setting.as_dict(), 'can_manage': True})


@require_http_methods(['GET', 'PATCH'])
def workspace_ai_settings(request, workspace_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    setting = _setting(workspace_id)
    provider_config = _safe_provider_config(setting)
    providers = {provider: values['has_api_key'] for provider, values in provider_config.items()}
    can_manage_access = membership.has_permission('manage_ai_access')
    can_manage_providers = membership.has_permission('manage_ai_providers')
    if request.method == 'GET':
        return JsonResponse({'settings': setting.as_dict(), 'can_manage': can_manage_access or can_manage_providers, 'providers': providers, 'provider_config': provider_config})
    if not can_manage_access and not can_manage_providers:
        return JsonResponse({'error': 'You do not have permission to manage Zuri settings.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)

    # Every field is validated into a local candidate value below - nothing is
    # written to `setting` until every check across the whole payload passes,
    # so a rejected request never leaves the row partially mutated in memory
    # only to silently discard those changes instead of persisting them.
    member_ids = {member.user_id for member in setting.workspace.memberships.all()}
    selected = [int(value) for value in payload.get('ai_user_ids', []) if str(value).isdigit()]
    if 'ai_user_ids' in payload and not can_manage_access:
        return JsonResponse({'error': 'You do not have permission to manage Zuri member access.'}, status=403)
    if not set(selected).issubset(member_ids):
        return JsonResponse({'error': 'AI access can only be granted to workspace members.'}, status=400)
    if 'ai_enabled' in payload and not can_manage_access:
        return JsonResponse({'error': 'You do not have permission to manage Zuri member access.'}, status=403)
    new_ai_enabled = bool(payload.get('ai_enabled', setting.ai_enabled))
    new_ai_user_ids = selected if 'ai_user_ids' in payload else setting.ai_user_ids

    provider_fields_touched = bool({'ai_model', 'ai_default_provider', 'provider_config', 'ai_enabled_providers'} & set(payload))
    if provider_fields_touched and not can_manage_providers:
        return JsonResponse({'error': 'You do not have permission to manage Zuri providers.'}, status=403)
    new_ai_model = str(payload.get('ai_model', setting.ai_model or os.environ.get('AI_MODEL', ''))).strip()[:120]
    new_ai_default_provider = payload.get('ai_default_provider', setting.ai_default_provider) if payload.get('ai_default_provider') in providers else setting.ai_default_provider

    stored_config = dict(setting.ai_provider_config or {})
    submitted_config = payload.get('provider_config', {})
    if submitted_config is not None and not isinstance(submitted_config, dict):
        return JsonResponse({'error': 'Provider configuration must be an object.'}, status=400)
    for provider, defaults in AI_PROVIDER_DEFAULTS.items():
        submitted = submitted_config.get(provider, {}) if isinstance(submitted_config, dict) else {}
        if not isinstance(submitted, dict):
            return JsonResponse({'error': f'Invalid {provider} configuration.'}, status=400)
        current = dict(stored_config.get(provider, {}))
        if 'base_url' in submitted:
            base_url = str(submitted.get('base_url', '')).strip().rstrip('/') or defaults['base_url']
            parsed = urlparse(base_url)
            if parsed.scheme != 'https' or not parsed.netloc:
                return JsonResponse({'error': f'{provider.title()} base URL must be a valid HTTPS address.'}, status=400)
            current['base_url'] = base_url[:500]
        if 'model' in submitted:
            current['model'] = str(submitted.get('model', '')).strip()[:120] or defaults['model']
        api_key = str(submitted.get('api_key', '')).strip()
        if api_key:
            current['api_key_encrypted'] = _encrypt_secret(api_key)
        elif submitted.get('clear_api_key'):
            current.pop('api_key_encrypted', None)
        stored_config[provider] = current

    # Recompute has_api_key against the *candidate* config (not the persisted
    # one) so a key submitted in this same request already counts when
    # deciding which providers may be enabled.
    candidate_setting_snapshot = WorkspaceSetting(ai_provider_config=stored_config, ai_default_provider=new_ai_default_provider, ai_model=new_ai_model)
    candidate_provider_config = _safe_provider_config(candidate_setting_snapshot)
    candidate_providers = {provider: values['has_api_key'] for provider, values in candidate_provider_config.items()}
    requested_enabled = payload.get('ai_enabled_providers', setting.ai_enabled_providers or [])
    missing_keys = [provider for provider in requested_enabled if provider in candidate_providers and not candidate_providers[provider]]
    if missing_keys:
        labels = ', '.join(provider.title() for provider in missing_keys)
        return JsonResponse({'error': f'Add and save an API key before enabling {labels}.'}, status=400)
    new_ai_enabled_providers = [provider for provider in requested_enabled if provider in candidate_providers and candidate_providers[provider]]

    # All validation passed - apply every field together and save once.
    setting.ai_enabled = new_ai_enabled
    setting.ai_user_ids = new_ai_user_ids
    setting.ai_model = new_ai_model
    setting.ai_default_provider = new_ai_default_provider
    setting.ai_provider_config = stored_config
    setting.ai_enabled_providers = new_ai_enabled_providers
    setting.save(update_fields=['ai_enabled', 'ai_user_ids', 'ai_model', 'ai_default_provider', 'ai_enabled_providers', 'ai_provider_config', 'updated_at'])
    provider_config = _safe_provider_config(setting)
    providers = {provider: values['has_api_key'] for provider, values in provider_config.items()}
    return JsonResponse({'settings': setting.as_dict(), 'can_manage': can_manage_access or can_manage_providers, 'providers': providers, 'provider_config': provider_config})


# The persona is fixed here, not per-provider, so switching between OpenAI/
# Claude/Kimi/DeepSeek (an admin setting) never changes who the assistant
# says it is - the underlying provider and model name are deliberately never
# disclosed to the end user.
AI_SYSTEM_PROMPT = (
    'You are Zuri, the workspace assistant for WorkSpace. If asked your name, you are Zuri. '
    'Never reveal or discuss which underlying AI provider or model powers you, even if asked directly - '
    'just say you are Zuri. Be concise, practical, and protect confidential information.'
)
# Bounds on the prior turns a client may replay. The transcript lives in the
# caller's browser, so treat it as untrusted input: keep it small enough that a
# long conversation cannot blow up token spend or the request body.
AI_HISTORY_MAX_TURNS = 20
AI_HISTORY_MAX_CHARS = 24000


def _ai_history(raw):
    """Normalise client-supplied prior turns into provider message dicts."""
    if not isinstance(raw, list):
        return []
    turns = []
    for entry in raw[-AI_HISTORY_MAX_TURNS:]:
        if not isinstance(entry, dict):
            continue
        role = entry.get('role')
        content = entry.get('content')
        if role not in {'user', 'assistant'} or not isinstance(content, str):
            continue
        content = content.strip()
        if content:
            turns.append({'role': role, 'content': content})
    # Trim from the front so the most recent context survives the budget.
    budget = AI_HISTORY_MAX_CHARS
    kept = []
    for turn in reversed(turns):
        budget -= len(turn['content'])
        if budget < 0:
            break
        kept.append(turn)
    kept.reverse()
    # Providers reject a leading assistant turn, which is what a mid-conversation
    # trim can easily leave behind.
    while kept and kept[0]['role'] == 'assistant':
        kept.pop(0)
    return kept


@require_http_methods(['POST'])
def workspace_ai_chat(request, workspace_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    setting = _setting(workspace_id)
    if not membership.has_permission('use_ai'):
        return JsonResponse({'error': 'You do not have permission to use Zuri.'}, status=403)
    if not setting.ai_enabled or (membership.role == 'member' and request.user.id not in (setting.ai_user_ids or [])):
        return JsonResponse({'error': 'Zuri has not been enabled for your account.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    message = str(payload.get('message', '')).strip()
    if not message or len(message) > 12000:
        return JsonResponse({'error': 'Enter a message up to 12,000 characters.'}, status=400)
    provider = str(payload.get('provider') or setting.ai_default_provider or 'openai').lower()
    if provider not in AI_PROVIDER_DEFAULTS:
        return JsonResponse({'error': 'Unknown AI provider.'}, status=400)
    if setting.ai_enabled_providers and provider not in setting.ai_enabled_providers:
        return JsonResponse({'error': 'That AI provider is not enabled by your workspace administrator.'}, status=403)
    history = _ai_history(payload.get('history'))
    provider_values = _provider_values(setting, provider)
    api_key = provider_values['api_key']
    if not api_key:
        return JsonResponse({'error': 'The company AI API key has not been configured yet.'}, status=503)
    model = provider_values['model']
    endpoint = _provider_endpoint(provider, provider_values['base_url'])
    try:
        privacy = PrivacyRegistry(workspace_id, request.user)
        message = privacy.protect(message)
        history = [
            {**turn, 'content': privacy.protect(turn['content'])}
            for turn in history
        ]
        snapshot = build_workspace_snapshot(workspace_id, request.user, privacy)
    except PrivacyBoundaryError as exc:
        return JsonResponse({'error': str(exc), 'code': 'privacy_boundary'}, status=400)
    system_prompt = f'{AI_SYSTEM_PROMPT}\n\n{action_instructions(snapshot)}'
    turns = history + [{'role': 'user', 'content': message}]
    if provider == 'claude':
        body = json.dumps({'model': model, 'max_tokens': 1200, 'system': system_prompt, 'messages': turns}).encode()
        req = urlrequest.Request(endpoint, data=body, headers={'x-api-key': api_key, 'anthropic-version': '2023-06-01', 'Content-Type': 'application/json'}, method='POST')
    else:
        body = json.dumps({'model': model, 'messages': [{'role': 'system', 'content': system_prompt}] + turns, 'temperature': 0.3}).encode()
        req = urlrequest.Request(endpoint, data=body, headers={'Authorization': f'Bearer {api_key}', 'Content-Type': 'application/json'}, method='POST')
    try:
        with urlrequest.urlopen(req, timeout=45) as response:
            result = json.loads(response.read().decode())
        answer = (result.get('content', [{}])[0].get('text', '') if provider == 'claude' else result.get('choices', [{}])[0].get('message', {}).get('content', '')).strip()
        parsed = parse_provider_response(answer, privacy)
        pending_action = None
        if parsed['action'] is not None:
            try:
                pending_action = create_action_proposal(parsed['action'], privacy, workspace_id, request.user)
            except (ActionValidationError, PrivacyBoundaryError) as exc:
                return JsonResponse({'answer': parsed['answer'], 'pending_action': None, 'action_error': str(exc)})
        return JsonResponse({
            'answer': parsed['answer'] or 'The assistant returned an empty response.',
            'pending_action': pending_action.as_dict() if pending_action else None,
        })
    except (HTTPError, URLError, TimeoutError, ValueError) as exc:
        return JsonResponse({'error': f'AI service unavailable: {exc}'}, status=502)


@require_http_methods(['POST'])
def workspace_ai_action(request, workspace_id, action_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    setting = _setting(workspace_id)
    if not membership.has_permission('use_ai'):
        return JsonResponse({'error': 'You do not have permission to use Zuri.'}, status=403)
    if not setting.ai_enabled or (membership.role == 'member' and request.user.id not in (setting.ai_user_ids or [])):
        return JsonResponse({'error': 'Zuri has not been enabled for your account.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    decision = str(payload.get('decision') or '').strip().lower()
    if decision not in {'confirm', 'cancel'}:
        return JsonResponse({'error': 'Decision must be confirm or cancel.'}, status=400)
    with transaction.atomic():
        action = (
            AiAction.objects
            .select_for_update()
            .filter(id=action_id, workspace_id=workspace_id, requested_by=request.user)
            .first()
        )
        if action is None:
            return JsonResponse({'error': 'Pending Zuri action was not found.'}, status=404)
        if action.status != 'pending':
            return JsonResponse({'action': action.as_dict()})
        if timezone.now() >= action.expires_at:
            action.status = 'expired'
            action.error = 'This action expired before it was confirmed.'
            action.resolved_at = timezone.now()
            action.save(update_fields=['status', 'error', 'resolved_at'])
            return JsonResponse({'error': action.error, 'action': action.as_dict()}, status=409)
        if decision == 'cancel':
            action.status = 'cancelled'
            action.resolved_at = timezone.now()
            action.save(update_fields=['status', 'resolved_at'])
            return JsonResponse({'action': action.as_dict()})
        try:
            action.result = execute_action(action, request.user)
            action.status = 'executed'
            action.error = ''
            action.resolved_at = timezone.now()
            action.save(update_fields=['result', 'status', 'error', 'resolved_at'])
        except (ActionExecutionError, ActionValidationError, PrivacyBoundaryError, ValueError) as exc:
            action.status = 'failed'
            action.error = str(exc)
            action.resolved_at = timezone.now()
            action.save(update_fields=['status', 'error', 'resolved_at'])
            return JsonResponse({'error': str(exc), 'action': action.as_dict()}, status=getattr(exc, 'status', 400))
    return JsonResponse({'action': action.as_dict()})


@require_http_methods(['GET', 'POST'])
def workspace_document_list(request, workspace_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    if request.method == 'GET':
        documents = list(WorkspaceDocument.objects.filter(workspace_id=workspace_id))
        shared_permissions = dict(WorkspaceDocumentShare.objects.filter(document__workspace_id=workspace_id, user=request.user).values_list('document_id', 'permission'))
        can_lead = membership.role in {'owner', 'manager'}
        payload = [{**document.as_dict(), 'permission': 'edit' if can_lead or document.created_by_id == request.user.id else shared_permissions.get(document.id, 'view')} for document in documents]
        return JsonResponse({'documents': payload})
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    title = str(payload.get('title', '')).strip()[:200] or 'Untitled document'
    kind = payload.get('kind', 'document') if payload.get('kind') in {'document', 'presentation', 'spreadsheet'} else 'document'
    document = WorkspaceDocument.objects.create(workspace_id=workspace_id, title=title, kind=kind, content=sanitize_document_content(payload.get('content') or {}), created_by=request.user)
    return JsonResponse({'document': {**document.as_dict(), 'permission': 'edit'}}, status=201)


@require_http_methods(['GET', 'PATCH', 'DELETE'])
def workspace_document_detail(request, workspace_id, document_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    permission = _document_permission(document, membership, request.user)
    if request.method == 'GET':
        return JsonResponse({'document': {**document.as_dict(), 'permission': permission}})
    if request.method == 'DELETE':
        if membership.role not in {'owner', 'manager'} and document.created_by_id != request.user.id:
            return JsonResponse({'error': 'Only the document owner or a workspace leader can delete it.'}, status=403)
        document.delete()
        return JsonResponse({'status': 'deleted'})
    if permission != 'edit':
        return JsonResponse({'error': 'Edit access is required.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    base_updated_at = str(payload.get('base_updated_at') or '').strip()
    if base_updated_at and base_updated_at != document.updated_at.isoformat():
        return JsonResponse({
            'error': 'Someone else saved this document while you were editing. Reload to see their changes.',
            'document': {**document.as_dict(), 'permission': permission},
        }, status=409)
    if 'title' in payload:
        document.title = str(payload['title']).strip()[:200] or document.title
    try:
        with transaction.atomic():
            if 'content' in payload and isinstance(payload['content'], dict):
                try:
                    json.dumps(payload['content'], allow_nan=False)
                except (TypeError, ValueError):
                    return JsonResponse({'error': 'Document content contains invalid values.'}, status=400)
                content = sanitize_document_content(payload['content'])
                if content != document.content:
                    _record_revision(document, request.user)
                document.content = content
            document.save(update_fields=['title', 'content', 'updated_at'])
    except Exception as exc:
        logger.exception('Document save failed for %s', document.id)
        return JsonResponse({'error': f'Document could not be saved: {exc}'}, status=500)
    return JsonResponse({'document': {**document.as_dict(), 'permission': permission}})


@require_http_methods(['GET'])
def workspace_document_revision_list(request, workspace_id, document_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    revisions = document.revisions.select_related('created_by')[:REVISION_KEEP]
    return JsonResponse({'revisions': [{
        'id': revision.id,
        'title': revision.title,
        'created_at': revision.created_at.isoformat(),
        'created_by': revision.created_by.get_full_name() or revision.created_by.email if revision.created_by else 'Unknown user',
    } for revision in revisions]})


@require_http_methods(['POST'])
def workspace_document_revision_restore(request, workspace_id, document_id, revision_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    permission = _document_permission(document, membership, request.user)
    if permission != 'edit':
        return JsonResponse({'error': 'Edit access is required.'}, status=403)
    revision = document.revisions.filter(id=revision_id).first()
    if not revision:
        return JsonResponse({'error': 'Version not found.'}, status=404)
    with transaction.atomic():
        # Snapshot what is on screen now so restoring is itself undoable.
        WorkspaceDocumentRevision.objects.create(document=document, created_by=request.user, title=document.title, content=document.content or {})
        document.title = revision.title
        document.content = sanitize_document_content(revision.content or {})
        document.save(update_fields=['title', 'content', 'updated_at'])
    return JsonResponse({'document': {**document.as_dict(), 'permission': permission}})


@require_http_methods(['GET'])
def workspace_document_export(request, workspace_id, document_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    if document.kind != 'spreadsheet':
        return JsonResponse({'error': 'Only spreadsheets can be exported here.'}, status=400)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_data in document.content.get('sheets', []):
        sheet = workbook.create_sheet(str(sheet_data.get('name') or 'Sheet'))
        for row in sheet_data.get('rows', []):
            sheet.append(row)
    if not workbook.worksheets:
        workbook.create_sheet('Sheet 1')
    stream = io.BytesIO()
    workbook.save(stream)
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{document.title[:80]}.xlsx"'
    return response


@require_http_methods(['POST'])
def workspace_spreadsheet_import(request, workspace_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    if membership.role not in {'owner', 'manager'}:
        return JsonResponse({'error': 'Only workspace leaders can import spreadsheets.'}, status=403)
    uploaded = request.FILES.get('file')
    if not uploaded:
        return JsonResponse({'error': 'Choose a CSV or XLSX file.'}, status=400)
    name = uploaded.name.lower()
    if name.endswith('.csv'):
        rows = list(csv.reader(io.TextIOWrapper(uploaded.file, encoding='utf-8-sig')))
        sheets = [{'name': 'Sheet 1', 'rows': rows}]
    elif name.endswith('.xlsx'):
        workbook = load_workbook(uploaded, read_only=True, data_only=False)
        sheets = [{'name': sheet.title, 'rows': [[cell.value if cell.value is not None else '' for cell in row] for row in sheet.iter_rows()]} for sheet in workbook.worksheets]
    else:
        return JsonResponse({'error': 'Only CSV and XLSX files are supported.'}, status=400)
    return JsonResponse({'sheets': sheets})


@require_http_methods(['GET', 'POST'])
def workspace_document_share_list(request, workspace_id, document_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    if request.method == 'GET':
        return JsonResponse({'shares': [share.as_dict() for share in document.shares.select_related('user', 'shared_by')]})
    if membership.role not in {'owner', 'manager'} and document.created_by_id != request.user.id:
        return JsonResponse({'error': 'Only the document owner or a workspace leader can share it.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    user_id, id_error = parse_int(payload.get('user_id'), 'Member')
    if id_error:
        return JsonResponse({'error': id_error}, status=400)
    permission = payload.get('permission', 'view')
    if permission not in {'view', 'comment', 'edit'}:
        return JsonResponse({'error': 'Permission must be view, comment, or edit.'}, status=400)
    member = Membership.objects.filter(workspace_id=workspace_id, user_id=user_id).select_related('user').first()
    if not member:
        return JsonResponse({'error': 'Choose a member of this workspace.'}, status=400)
    previous = WorkspaceDocumentShare.objects.filter(document=document, user=member.user).values_list('permission', flat=True).first()
    share, _ = WorkspaceDocumentShare.objects.update_or_create(document=document, user=member.user, defaults={'permission': permission, 'shared_by': request.user})
    if member.user_id != request.user.id and previous != permission:
        from .views import create_notification
        actor_name = request.user.get_full_name() or request.user.email
        create_notification(
            workspace_id, member.user, 'document_shared',
            f'{actor_name} shared "{document.title}" with you',
            f'You have {permission} access.' if previous is None else f'Your access is now {permission}.',
            target_type='document', target_id=document.id,
        )
    return JsonResponse({'share': share.as_dict()}, status=201)


@require_http_methods(['DELETE'])
def workspace_document_share_detail(request, workspace_id, document_id, share_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    if membership.role not in {'owner', 'manager'} and document.created_by_id != request.user.id:
        return JsonResponse({'error': 'Only the document owner or a workspace leader can change sharing.'}, status=403)
    share = document.shares.filter(id=share_id).first()
    if not share:
        return JsonResponse({'error': 'Share not found.'}, status=404)
    share.delete()
    return JsonResponse({'status': 'deleted'})


@require_http_methods(['GET', 'POST'])
def workspace_document_comment_list(request, workspace_id, document_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    document = WorkspaceDocument.objects.filter(workspace_id=workspace_id, id=document_id).first()
    if not document:
        return JsonResponse({'error': 'Document not found.'}, status=404)
    if request.method == 'GET':
        return JsonResponse({'comments': [comment.as_dict() for comment in document.comments.select_related('author')]})
    if _document_permission(document, membership, request.user) not in {'comment', 'edit'}:
        return JsonResponse({'error': 'Comment access is required.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    body = str(payload.get('body', '')).strip()
    if not body or len(body) > 4000:
        return JsonResponse({'error': 'Comment must be between 1 and 4,000 characters.'}, status=400)
    parent_id, id_error = parse_int(payload.get('parent_id'), 'Parent comment')
    if id_error:
        return JsonResponse({'error': id_error}, status=400)
    parent = document.comments.filter(id=parent_id).first() if parent_id else None
    comment = WorkspaceDocumentComment.objects.create(document=document, author=request.user, parent=parent, body=body, anchor=payload.get('anchor') if isinstance(payload.get('anchor'), dict) else {})
    from .views import create_notification, notify_mentions
    # The document owner and everyone already in the thread hear about a new
    # reply. Mentions exclude them so a comment cannot land twice for one
    # recipient: once as a reply and once as an @mention.
    recipient_ids = set(document.comments.exclude(author=request.user).values_list('author_id', flat=True))
    if document.created_by_id:
        recipient_ids.add(document.created_by_id)
    recipient_ids.discard(request.user.id)
    recipient_ids.discard(None)
    if recipient_ids:
        actor_name = request.user.get_full_name() or request.user.email
        # Scoped through Membership so someone who has since left the workspace
        # does not keep receiving notifications for a document they lost access to.
        members = Membership.objects.filter(workspace_id=workspace_id, user_id__in=recipient_ids).select_related('user')
        for member in members:
            create_notification(
                workspace_id, member.user, 'document_comment',
                f'{actor_name} commented on "{document.title}"',
                body[:120],
                target_type='document', target_id=document.id,
            )
    notify_mentions(workspace_id, request.user, body, 'document', document.id, exclude_user_ids=recipient_ids)
    return JsonResponse({'comment': comment.as_dict()}, status=201)


@require_http_methods(['PATCH'])
def workspace_document_comment_detail(request, workspace_id, document_id, comment_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    comment = WorkspaceDocumentComment.objects.filter(document_id=document_id, document__workspace_id=workspace_id, id=comment_id).first()
    if not comment:
        return JsonResponse({'error': 'Comment not found.'}, status=404)
    if membership.role not in {'owner', 'manager'} and comment.document.created_by_id != request.user.id and comment.author_id != request.user.id:
        return JsonResponse({'error': 'Only the comment author, document owner, or a workspace leader can update it.'}, status=403)
    try:
        payload = json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Request body must be valid JSON.'}, status=400)
    # Resolution is the only thing this endpoint changes, so an explicit value is
    # required. Defaulting a missing field to true meant an empty PATCH silently
    # resolved the comment.
    if 'resolved' not in payload:
        return JsonResponse({'error': 'resolved is required.'}, status=400)
    if payload['resolved']:
        from django.utils import timezone
        comment.resolved_at = timezone.now()
        comment.resolved_by = request.user
    else:
        comment.resolved_at = None
        comment.resolved_by = None
    comment.save(update_fields=['resolved_at', 'resolved_by'])
    return JsonResponse({'comment': comment.as_dict()})


UPLOAD_MAX_BYTES = getattr(django_settings, 'WORKSPACE_UPLOAD_MAX_BYTES', 25 * 1024 * 1024)

# Deliberately excludes anything a browser or the OS will execute.
UPLOAD_ALLOWED_EXTENSIONS = {
    '.csv', '.doc', '.docx', '.gif', '.jpeg', '.jpg', '.json', '.md', '.mp3', '.mp4', '.odp', '.ods',
    '.odt', '.pdf', '.png', '.ppt', '.pptx', '.rtf', '.txt', '.wav', '.webm', '.webp', '.xls',
    '.xlsx', '.xml', '.zip',
}


def _reject_upload(uploaded):
    """Return an error response when an upload fails size or type checks."""
    if uploaded.size > UPLOAD_MAX_BYTES:
        return JsonResponse({'error': f'Files must be {UPLOAD_MAX_BYTES // (1024 * 1024)} MB or smaller.'}, status=400)
    extension = os.path.splitext(uploaded.name or '')[1].lower()
    if extension not in UPLOAD_ALLOWED_EXTENSIONS:
        return JsonResponse({'error': f'{extension or "That file type"} is not an allowed file type.'}, status=400)
    return None


@require_http_methods(['GET', 'POST'])
def workspace_file_list(request, workspace_id):
    _, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    if request.method == 'GET':
        return JsonResponse({'files': [item.as_dict() for item in WorkspaceFile.objects.filter(workspace_id=workspace_id)]})
    uploaded = request.FILES.get('file')
    if not uploaded:
        return JsonResponse({'error': 'Choose a file to upload.'}, status=400)
    invalid = _reject_upload(uploaded)
    if invalid:
        return invalid
    # Trust the extension we validated, not the client-supplied content type.
    mime_type = mimetypes.guess_type(uploaded.name)[0] or uploaded.content_type or ''
    item = WorkspaceFile.objects.create(workspace_id=workspace_id, file=uploaded, original_name=uploaded.name[:255], mime_type=mime_type[:160], size=uploaded.size, uploaded_by=request.user)
    cloudinary_url = os.environ.get('CLOUDINARY_URL', '').strip()
    if cloudinary_url:
        try:
            import cloudinary.uploader
            uploaded.seek(0)
            result = cloudinary.uploader.upload(uploaded, resource_type='auto', folder=f'workspace/{workspace_id}')
            item.cloudinary_url = result.get('secure_url', '')
            item.cloudinary_public_id = result.get('public_id', '')
            item.save(update_fields=['cloudinary_url', 'cloudinary_public_id'])
        except Exception:
            logger.exception('Cloudinary upload failed for workspace file %s', item.id)
    return JsonResponse({'file': item.as_dict()}, status=201)


@require_http_methods(['GET'])
def workspace_file_download(request, file_id):
    if not request.user.is_authenticated:
        return JsonResponse({'error': 'Authentication is required.'}, status=401)
    item = WorkspaceFile.objects.filter(id=file_id, workspace__members=request.user).first()
    if not item:
        return JsonResponse({'error': 'File not found.'}, status=404)
    if not item.file:
        # Stored only in Cloudinary. Membership is already checked above, so the
        # redirect is the one place the public URL is handed out.
        if item.cloudinary_url:
            return HttpResponseRedirect(item.cloudinary_url)
        return JsonResponse({'error': 'File not found.'}, status=404)
    try:
        stored = item.file.open('rb')
    except (FileNotFoundError, OSError):
        # The row survives a redeploy but the file it names does not: uploads live
        # on the container's disk unless Cloudinary is configured.
        return JsonResponse({'error': 'File is no longer stored. Upload it again.'}, status=404)
    return stored_file_response(request, stored, item.original_name)


@require_http_methods(['DELETE'])
def workspace_file_detail(request, workspace_id, file_id):
    membership, error = require_workspace_member(request, workspace_id)
    if error:
        return error
    item = WorkspaceFile.objects.filter(workspace_id=workspace_id, id=file_id).first()
    if not item:
        return JsonResponse({'error': 'File not found.'}, status=404)
    if membership.role not in {'owner', 'manager'} and item.uploaded_by_id != request.user.id:
        return JsonResponse({'error': 'Only the uploader or a workspace leader can delete this file.'}, status=403)
    if item.cloudinary_public_id:
        try:
            import cloudinary.uploader
            cloudinary.uploader.destroy(item.cloudinary_public_id, invalidate=True)
        except Exception:
            logger.exception('Cloudinary delete failed for workspace file %s', item.id)
    if item.file:
        try:
            item.file.delete(save=False)
        except (OSError, ValueError):
            logger.warning('Stored file could not be removed for workspace file %s', item.id, exc_info=True)
    item.delete()
    return JsonResponse({'status': 'deleted'})
